from pathlib import Path

from settings.paths import CHARTS_DIR, DASHBOARD_DIR

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side
)
from openpyxl.drawing.image import Image


def load_table(engine, table_name):

    return pd.read_sql(
        f"SELECT * FROM {table_name}",
        engine
    )


def add_chart_image(
    worksheet,
    image_path,
    cell,
    width=500,
    height=250
):

    if not Path(image_path).exists():
        print(f"Missing chart: {image_path}")
        return

    img = Image(image_path)

    img.width = width
    img.height = height

    worksheet.add_image(
        img,
        cell
    )


def excel_dashboard(engine):

    dashboard_dir = DASHBOARD_DIR

    dashboard_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    output_file = (
        dashboard_dir /
        "trading_dashboard.xlsx"
    )

    wb = Workbook()

    dashboard = wb.active
    dashboard.title = "Dashboard"

    dashboard.sheet_view.showGridLines = False

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    table_names = [
        "avg_win_loss",
        "best_worst_trade",
        "daily_pnl",
        "drawdown_analysis",
        "equity_curve",
        "longest_winning_streak",
        "monthly_pnl",
        "monthly_rank",
        "pnl_by_symbol",
        "profit_factor",
        "session_analysis",
        "time_day_analysis",
        "total_pnl",
        "trade_counts",
        "weekday_analysis",
        "win_rate"
    ]

    tables = {}

    # ====================================
    # RAW DATA SHEETS
    # ====================================

    for table_name in table_names:

        df = load_table(
            engine,
            table_name
        )

        tables[table_name] = df

        ws = wb.create_sheet(
            title=table_name
        )

        ws.append(
            list(df.columns)
        )

        for row in df.itertuples(index=False):
            ws.append(list(row))

    # ====================================
    # TITLE
    # ====================================

    dashboard.merge_cells("A1:J1")

    dashboard["A1"] = (
        "Trading Performance Analytics Dashboard"
    )

    dashboard["A1"].font = Font(
        bold=True,
        size=16
    )

    dashboard["A1"].alignment = Alignment(
        horizontal="center"
    )

    # ====================================
    # KPI LABELS
    # ====================================

    dashboard["A3"] = "Total PnL"
    dashboard["C3"] = "Win Rate"
    dashboard["E3"] = "Profit Factor"
    dashboard["G3"] = "Total Trades"
    dashboard["I3"] = "Winning Streak"

    dashboard["A7"] = "Best Trade"
    dashboard["C7"] = "Worst Trade"
    dashboard["E7"] = "Average Win"
    dashboard["G7"] = "Average Loss"

    for cell in [
        "A3", "C3", "E3", "G3", "I3",
        "A7", "C7", "E7", "G7"
    ]:

        dashboard[cell].fill = header_fill

        dashboard[cell].font = Font(
            color="FFFFFF",
            bold=True
        )

        dashboard[cell].alignment = Alignment(
            horizontal="center"
        )

        dashboard[cell].border = thin_border

    # ====================================
    # KPI DATA
    # ====================================

    total_pnl = (
        tables["total_pnl"]["total_pnl"].iloc[0]
    )

    win_rate = (
        tables["win_rate"]["win_rate"].iloc[0]
    )

    profit_factor = (
        tables["profit_factor"]["profit_factor"].iloc[0]
    )

    total_trades = (
        tables["trade_counts"]["total_trades"].iloc[0]
    )

    best_trade = (
        tables["best_worst_trade"]["best_trade"].iloc[0]
    )

    worst_trade = (
        tables["best_worst_trade"]["worst_trade"].iloc[0]
    )

    avg_win = (
        tables["avg_win_loss"]["average_win"].iloc[0]
    )

    avg_loss = (
        tables["avg_win_loss"]["average_loss"].iloc[0]
    )

    streak = (
        tables["longest_winning_streak"]
        ["longest_winning_streak"]
        .iloc[0]
    )

    dashboard["A4"] = round(total_pnl, 2)
    dashboard["C4"] = round(win_rate, 2)
    dashboard["E4"] = round(profit_factor, 2)
    dashboard["G4"] = int(total_trades)
    dashboard["I4"] = int(streak)

    dashboard["A8"] = round(best_trade, 2)
    dashboard["C8"] = round(worst_trade, 2)
    dashboard["E8"] = round(avg_win, 2)
    dashboard["G8"] = round(avg_loss, 2)

    for cell in [
        "A4", "C4", "E4", "G4", "I4",
        "A8", "C8", "E8", "G8"
    ]:

        dashboard[cell].alignment = Alignment(
            horizontal="center"
        )

        dashboard[cell].border = thin_border

        dashboard[cell].font = Font(
            bold=True
        )

    if total_pnl < 0:

        dashboard["A4"].font = Font(
            color="FF0000",
            bold=True
        )

    else:

        dashboard["A4"].font = Font(
            color="008000",
            bold=True
        )

    # ====================================
    # CHART IMAGES
    # ====================================

    add_chart_image(
        dashboard,
        CHARTS_DIR / "equity_curve.png",
        "A12"
    )

    add_chart_image(
        dashboard,
        CHARTS_DIR / "drawdown.png",
        "F12"
    )

    add_chart_image(
        dashboard,
        CHARTS_DIR / "daily_pnl.png",
        "A27"
    )

    add_chart_image(
        dashboard,
        CHARTS_DIR / "pnl_by_symbol.png",
        "F27"
    )

    add_chart_image(
        dashboard,
        CHARTS_DIR / "session_analysis.png",
        "A42"
    )

    add_chart_image(
        dashboard,
        CHARTS_DIR / "weekday_analysis.png",
        "F42"
    )

    # ====================================
    # LAYOUT
    # ====================================

    for col in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
        dashboard.column_dimensions[col].width = 14

    # ====================================
    # HIDE DATA SHEETS
    # ====================================

    for sheet_name in wb.sheetnames:

        if sheet_name != "Dashboard":

            wb[sheet_name].sheet_state = "hidden"

    wb.save(output_file)

    print(
        f"Dashboard created: {output_file}"
    )
#EOF